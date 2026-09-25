"""Disk-backed workbook data. Never materialize an entire workbook in RAM."""
import csv
import codecs
import datetime as dt
import json
import math
import sqlite3
from functools import lru_cache
from pathlib import Path
from xml.etree.ElementTree import tostring, fromstring
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.text import Text
from openpyxl.reader.excel import ExcelReader
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from openpyxl.xml.constants import SHARED_STRINGS, SHEET_MAIN_NS
from openpyxl.xml.functions import iterparse

MAX_ROWS, MAX_COLS = 1048576, 16384

class Cancelled(Exception):
    pass

class ClosingConnection(sqlite3.Connection):
    def __exit__(self, *args):
        try:
            return super().__exit__(*args)
        finally:
            self.close()

def connect(path):
    db = sqlite3.connect(path, timeout=30, factory=ClosingConnection)
    db.execute('PRAGMA cache_size=-8192')
    db.execute('PRAGMA temp_store=FILE')
    return db

def init_db(path):
    with connect(path) as db:
        db.executescript('''
        PRAGMA journal_mode=WAL;
        CREATE TABLE cells (sheet INTEGER, row INTEGER, col INTEGER, data TEXT, text TEXT,
                            PRIMARY KEY(sheet,row,col)) WITHOUT ROWID;
        CREATE TABLE strings (id INTEGER PRIMARY KEY, value TEXT);
        ''')

class DiskStrings:
    def __init__(self, db):
        self.db = db
        self.cached_lookup = lru_cache(maxsize=256)(self.lookup)

    def lookup(self, index):
        value = self.db.execute('SELECT value FROM strings WHERE id=?', (index,)).fetchone()
        if value is None:
            raise ValueError('Cadena compartida no encontrada en el XLSX.')
        return value[0]

    def __getitem__(self, index):
        return self.cached_lookup(index)

class StreamingWorksheet(ReadOnlyWorksheet):
    def _get_size(self):
        # We count actual rows below. Without <dimension>, openpyxl's preliminary
        # dimension scan retains the whole sheet tree before its streaming pass.
        pass

class DiskReader(ExcelReader):
    """openpyxl 3.1.5 reader with its shared-string array replaced by SQLite."""
    def __init__(self, filename, db, report, check):
        self.db, self.report, self.check = db, report, check
        super().__init__(filename, read_only=True, data_only=False, keep_links=False)

    def read_strings(self):
        part = self.package.find(SHARED_STRINGS)
        if part is None:
            return
        with self.archive.open(part.PartName[1:]) as source:
            root = None
            index = 0
            for event, node in iterparse(source, events=('start', 'end')):
                if root is None:
                    root = node
                if event == 'end' and node.tag == '{%s}si' % SHEET_MAIN_NS:
                    value = Text.from_tree(node).content.replace('x005F_', '')
                    self.db.execute('INSERT INTO strings VALUES (?,?)', (index, value))
                    index += 1
                    root.clear()
                    if index % 1000 == 0:
                        self.check()
                        self.report(f'Leyendo textos: {index:,}', index)
                        self.db.commit()
        self.db.commit()
        self.shared_strings = DiskStrings(self.db)

    def read_worksheets(self):
        for sheet, rel in self.parser.find_sheets():
            if rel.target not in self.valid_files:
                continue
            if 'chartsheet' in rel.Type:
                self.read_chartsheet(sheet, rel)
                continue
            ws = StreamingWorksheet(self.wb, sheet.name, rel.target, self.shared_strings)
            ws.sheet_state = sheet.state
            self.wb._sheets.append(ws)

def encode_cell(cell):
    value = cell.value
    if value is None:
        return None
    if isinstance(value, ArrayFormula):
        result = {'t': 'n', 'f': (value.text or '=')[1:], 'F': value.ref}
    elif isinstance(value, DataTableFormula):
        result = {'t': 'n', 'tableFormula': dict(value)}
    elif cell.data_type == 'f':
        result = {'t': 'n', 'f': value[1:]}
    elif isinstance(value, bool):
        result = {'t': 'b', 'v': value}
    elif isinstance(value, dt.timedelta):
        result = {'t': 'n', 'v': value.total_seconds() / 86400}
    elif isinstance(value, (dt.datetime, dt.date, dt.time)):
        result = {'t': 'd', 'v': value.isoformat(), 'dateKind': type(value).__name__}
    elif isinstance(value, (int, float)):
        result = {'t': 'n', 'v': value}
    else:
        result = {'t': 'e' if cell.data_type == 'e' else 's', 'v': str(value)}
    if cell.number_format != 'General':
        result['z'] = cell.number_format
    return result

def search_text(cell):
    return str(cell.get('v', '=' + cell.get('f', ''))).casefold() if cell else ''

def names(container):
    return [tostring(name.to_tree(), encoding='unicode') for name in container.values()]

def import_file(source, db_path, report=lambda *_: None, check=lambda: None):
    init_db(db_path)
    db = connect(db_path)
    workbook = None
    reader = None
    try:
        if Path(source).suffix.lower() == '.xlsx':
            reader = DiskReader(source, db, report, check)
            reader.read()
            workbook = reader.wb
            meta = {'epoch': workbook.epoch.isoformat(), 'names': names(workbook.defined_names), 'sheets': []}
            for sid, ws in enumerate(workbook.worksheets):
                ws.reset_dimensions()  # Do not trust incorrect or inflated producer dimensions.
                rows, cols, batch = 0, 0, []
                for ri, row in enumerate(ws.iter_rows(), 1):
                    if ri > MAX_ROWS:
                        raise ValueError('La hoja supera las 1.048.576 filas de Excel.')
                    for ci, cell in enumerate(row, 1):
                        item = encode_cell(cell)
                        if item is None:
                            continue
                        if ci > MAX_COLS:
                            raise ValueError('La hoja supera las 16.384 columnas de Excel.')
                        rows, cols = max(rows, ri), max(cols, ci)
                        batch.append((sid, ri, ci, json.dumps(item, ensure_ascii=False), search_text(item)))
                        if len(batch) >= 1000:
                            db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
                            batch.clear()
                    if ri % 200 == 0:
                        check()
                        db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
                        batch.clear()
                        db.commit()
                        report(f'{ws.title}: {ri:,} filas leídas', ri)
                db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
                db.commit()
                meta['sheets'].append({'name': ws.title, 'rows': max(rows, 1), 'cols': max(cols, 1), 'names': names(ws.defined_names)})
        else:
            meta = {'sheets': []}
            # utf-8-sig handles Excel CSV with BOM; deterministic Latin-1 fallback.
            with open(source, 'rb') as sample:
                prefix = sample.read(65536)
            try:
                codecs.getincrementaldecoder('utf-8-sig')().decode(prefix, final=False); encoding = 'utf-8-sig'
            except UnicodeDecodeError:
                encoding = 'latin-1'
            if prefix.startswith((b'\xff\xfe', b'\xfe\xff')):
                encoding = 'utf-16'
            sample = prefix.decode(encoding, errors='replace')
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            rows, cols, batch = 0, 0, []
            csv.field_size_limit(1024 * 1024)
            with open(source, encoding=encoding, newline='') as file:
                for ri, row in enumerate(csv.reader(file, dialect), 1):
                    if ri > MAX_ROWS or len(row) > MAX_COLS:
                        raise ValueError('El CSV supera las dimensiones de Excel; divídelo en hojas.')
                    rows, cols = ri, max(cols, len(row))
                    for ci, value in enumerate(row, 1):
                        if len(value) > 32767:
                            raise ValueError(f'Celda {ri},{ci}: supera los 32.767 caracteres que admite Excel.')
                        if value:
                            item = {'t': 's', 'v': value}
                            batch.append((0, ri, ci, json.dumps(item, ensure_ascii=False), value.casefold()))
                            if len(batch) >= 1000:
                                db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
                                batch.clear()
                    if ri % 200 == 0:
                        check()
                        db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
                        batch.clear(); db.commit(); report(f'CSV: {ri:,} filas leídas', ri)
            db.executemany('INSERT INTO cells VALUES (?,?,?,?,?)', batch)
            db.commit()
            meta['sheets'].append({'name': 'Hoja 1', 'rows': max(1, rows), 'cols': max(1, cols), 'names': []})
        if not meta['sheets']:
            raise ValueError('El archivo no contiene hojas.')
        return meta
    finally:
        if workbook is not None:
            workbook.close()
        elif reader is not None:
            reader.archive.close()
        db.close()

def get_page(db_path, meta, sid=0, offset=0, col=1, query=''):
    info = meta['sheets'][sid]
    offset = max(0, offset); col = max(1, min(col, info['cols']))
    with connect(db_path) as db:
        if query:
            rows = [x[0] for x in db.execute('SELECT DISTINCT row FROM cells WHERE sheet=? AND instr(text,?)>0 ORDER BY row LIMIT 101 OFFSET ?', (sid, query.casefold(), offset))]
            more, rows = len(rows) > 100, rows[:100]
        else:
            rows = list(range(offset + 1, min(offset + 100, info['rows']) + 1))
            more = offset + 100 < info['rows']
        cells = {}
        if rows:
            params = [sid, *rows, col, min(col + 23, info['cols'])]
            placeholders = ','.join('?' for _ in rows)
            for r, c, data in db.execute(f'SELECT row,col,data FROM cells WHERE sheet=? AND row IN ({placeholders}) AND col BETWEEN ? AND ? ORDER BY row,col', params):
                cells[f'{r}:{c}'] = json.loads(data)
    return {'rows': rows, 'col': col, 'cols': min(24, info['cols'] - col + 1), 'cells': cells, 'more': more, 'offset': offset}

def validate_cell(cell):
    if cell is None:
        return
    if not isinstance(cell, dict) or cell.get('t') not in ('s', 'n', 'b', 'd', 'e'):
        raise ValueError('Tipo de celda inválido.')
    if len(str(cell.get('v', ''))) > 32767 or len(str(cell.get('f', ''))) > 8192:
        raise ValueError('El contenido supera el tamaño de celda o fórmula de Excel.')
    if 'f' in cell:
        if not isinstance(cell['f'], str) or not cell['f']:
            raise ValueError('Fórmula inválida.')
    elif cell['t'] == 'n' and (isinstance(cell.get('v'), bool) or not isinstance(cell.get('v'), (int, float)) or not math.isfinite(cell['v'])):
        raise ValueError('Número inválido.')

def edit_cells(db_path, meta, sid, changes):
    if not isinstance(changes, list) or not 1 <= len(changes) <= 10000:
        raise ValueError('Pega hasta 10.000 celdas por operación.')
    info = meta['sheets'][sid]
    seen = set()
    for item in changes:
        r, c = item['r'], item['c']
        if type(r) is not int or type(c) is not int or not (1 <= r <= info['rows'] and 1 <= c <= info['cols']) or (r, c) in seen:
            raise ValueError('Celda fuera de la hoja o repetida.')
        seen.add((r, c)); validate_cell(item['cell'])
    before = []
    with connect(db_path) as db:
        for item in changes:
            r, c, cell = item['r'], item['c'], item['cell']
            old = db.execute('SELECT data FROM cells WHERE sheet=? AND row=? AND col=?', (sid, r, c)).fetchone()
            old_cell = json.loads(old[0]) if old else None
            before.append({'r': r, 'c': c, 'cell': old_cell})
            if cell and old_cell and old_cell.get('z') and cell['t'] in ('n', 'd'):
                cell = {**cell, 'z': cell.get('z', old_cell['z'])}
            if cell is None:
                db.execute('DELETE FROM cells WHERE sheet=? AND row=? AND col=?', (sid, r, c))
            else:
                db.execute('INSERT OR REPLACE INTO cells VALUES (?,?,?,?,?)', (sid, r, c, json.dumps(cell, ensure_ascii=False), search_text(cell)))
    return before

def decode_value(item):
    if item.get('F'):
        return ArrayFormula(ref=item['F'], text='=' + item['f'])
    if item.get('tableFormula'):
        return DataTableFormula(**item['tableFormula'])
    if 'f' in item:
        return '=' + item['f']
    value = item.get('v')
    if item['t'] == 'd':
        return getattr(dt, item.get('dateKind', 'datetime')).fromisoformat(value)
    return value

def row_stream(db, sid, count):
    cursor = iter(db.execute('SELECT row,col,data FROM cells WHERE sheet=? ORDER BY row,col', (sid,)))
    current = next(cursor, None)
    for ri in range(1, count + 1):
        row = {}
        while current and current[0] == ri:
            row[current[1]] = json.loads(current[2]); current = next(cursor, None)
        yield ri, row

def export_file(db_path, meta, destination, kind='xlsx', sid=0, report=lambda *_: None, check=lambda: None):
    db = connect(db_path)
    wb = None
    try:
        db.execute('BEGIN')  # Consistent export snapshot.
        if kind == 'csv':
            info = meta['sheets'][sid]
            with open(destination, 'w', encoding='utf-8-sig', newline='') as out:
                writer = csv.writer(out)
                for ri, row in row_stream(db, sid, info['rows']):
                    values = []
                    for ci in range(1, info['cols'] + 1):
                        cell = row.get(ci)
                        value = decode_value(cell) if cell else ''
                        if cell and (cell.get('f') or cell['t'] == 's') and str(value).startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
                            value = "'" + str(value)
                        values.append(value)
                    writer.writerow(values)
                    if ri % 200 == 0:
                        check(); report(f'Exportando CSV: {ri:,} filas', ri)
        else:
            wb = Workbook(write_only=True)
            if meta.get('epoch'):
                wb.epoch = dt.datetime.fromisoformat(meta['epoch'])
            wb.calculation = CalcProperties(fullCalcOnLoad=True, forceFullCalc=True)
            for xml in meta.get('names', []):
                wb.defined_names.add(DefinedName.from_tree(fromstring(xml)))
            for sheet_id, info in enumerate(meta['sheets']):
                ws = wb.create_sheet(info['name'])
                for xml in info.get('names', []):
                    ws.defined_names.add(DefinedName.from_tree(fromstring(xml)))
                for ri, row in row_stream(db, sheet_id, info['rows']):
                    values = [None] * max(row, default=0)
                    for ci, data in row.items():
                        cell = WriteOnlyCell(ws, value=decode_value(data))
                        if data['t'] == 's' and 'f' not in data:
                            cell.data_type = 's'  # Formula-looking imported text stays text.
                        if data['t'] == 'e':
                            cell.data_type = 'e'
                        if data.get('z'):
                            cell.number_format = data['z']
                        values[ci - 1] = cell
                    ws.append(values)
                    if ri % 200 == 0:
                        check(); report(f'Exportando {info["name"]}: {ri:,} filas', ri)
            check(); report('Comprimiendo la copia Excel…', 0)
            wb.save(destination)
    finally:
        db.close()
        if wb is not None:
            for ws in wb.worksheets:
                if not ws.closed:
                    ws.close()
                if ws._writer and Path(ws._writer.out).exists():
                    ws._writer.cleanup()
