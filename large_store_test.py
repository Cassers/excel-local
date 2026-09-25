import datetime as dt
import json
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
import large_store as store

class LargeStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = self.root / 'data.sqlite'

    def tearDown(self):
        self.temp.cleanup()

    def fixture(self):
        wb = Workbook(); ws = wb.active; ws.title = 'Datos'
        ws.append(['Nombre', 'Número', 'Código', 'Fecha', 'Fórmula', 'Literal'])
        ws.append(['Árbol', 12.5, '00123', dt.datetime(2026, 9, 25), '=B2*2', '=literal'])
        ws['F2'].data_type = 's'; ws['B2'].number_format = '0.00'
        for i in range(3, 225): ws.append([f'Fila {i}', i])
        ws['AD224'] = 'columna 30'
        wb.create_sheet('Segunda')['A1'] = True
        wb.defined_names.add(DefinedName('Importe', attr_text='Datos!$B$2'))
        source = self.root / 'input.xlsx'; wb.save(source)
        return source

    def test_import_page_edit_export_roundtrip(self):
        meta = store.import_file(self.fixture(), self.db)
        self.assertEqual(meta['sheets'][0]['rows'], 224)
        self.assertEqual(meta['sheets'][0]['cols'], 30)
        first = store.get_page(self.db, meta)
        self.assertEqual(len(first['rows']), 100)
        self.assertEqual(first['cells']['2:3']['v'], '00123')
        self.assertEqual(first['cells']['2:5']['f'], 'B2*2')
        self.assertEqual(store.get_page(self.db, meta, query='ÁRBOL')['rows'], [2])
        self.assertEqual(store.get_page(self.db, meta, offset=200, col=25)['cells']['224:30']['v'], 'columna 30')
        old = store.edit_cells(self.db, meta, 0, [{'r': 2, 'c': 2, 'cell': {'t': 'n', 'v': 25}}])
        self.assertEqual(old[0]['cell']['v'], 12.5)
        out = self.root / 'output.xlsx'; store.export_file(self.db, meta, out)
        wb = load_workbook(out); ws = wb['Datos']
        self.assertEqual(ws['B2'].value, 25)
        self.assertEqual(ws['B2'].number_format, '0.00')
        self.assertEqual(ws['C2'].value, '00123')
        self.assertEqual(ws['D2'].value, dt.datetime(2026, 9, 25))
        self.assertEqual(ws['E2'].value, '=B2*2')
        self.assertEqual(ws['F2'].data_type, 's')
        self.assertEqual(wb.defined_names['Importe'].attr_text, 'Datos!$B$2')
        self.assertTrue(wb['Segunda']['A1'].value); wb.close()

    def test_invalid_edit_is_atomic(self):
        meta = store.import_file(self.fixture(), self.db)
        with self.assertRaises(ValueError):
            store.edit_cells(self.db, meta, 0, [{'r': 2, 'c': 2, 'cell': {'t': 'n', 'v': 999}}, {'r': 999, 'c': 1, 'cell': None}])
        self.assertEqual(store.get_page(self.db, meta)['cells']['2:2']['v'], 12.5)

    def test_csv_preserves_text_and_sanitizes_export(self):
        source = self.root / 'input.csv'; source.write_text('Nombre;Código;Literal\nCafé;00123;=1+1\n', encoding='utf-8-sig')
        meta = store.import_file(source, self.db)
        self.assertEqual(store.get_page(self.db, meta)['cells']['2:2']['v'], '00123')
        out = self.root / 'out.csv'; store.export_file(self.db, meta, out, 'csv')
        self.assertIn("'=1+1", out.read_text(encoding='utf-8-sig'))

    def test_cancel_import(self):
        source = self.root / 'input.csv'; source.write_text('hello\n' * 1000)
        def cancel(): raise store.Cancelled()
        with self.assertRaises(store.Cancelled): store.import_file(source, self.db, check=cancel)

    def test_shared_strings_are_stored_on_disk(self):
        source = self.fixture(); modified = self.root / 'shared.xlsx'
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        with ZipFile(source) as z, ZipFile(modified, 'w', ZIP_DEFLATED) as out:
            for item in z.infolist():
                data = z.read(item.filename)
                if item.filename == '[Content_Types].xml':
                    root = ET.fromstring(data)
                    ET.SubElement(root, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', PartName='/xl/sharedStrings.xml', ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
                    data = ET.tostring(root)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    root = ET.fromstring(data)
                    for c in root.iter(ns + 'c'):
                        if c.get('r') == 'A2':
                            c.clear(); c.set('r', 'A2'); c.set('t', 's'); ET.SubElement(c, ns + 'v').text = '0'
                    data = ET.tostring(root)
                out.writestr(item, data)
            out.writestr('xl/sharedStrings.xml', '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>Texto compartido</t></si></sst>')
        meta = store.import_file(modified, self.db)
        self.assertEqual(store.get_page(self.db, meta)['cells']['2:1']['v'], 'Texto compartido')
        with store.connect(self.db) as db: self.assertEqual(db.execute('SELECT count(*) FROM strings').fetchone()[0], 1)

if __name__ == '__main__': unittest.main()
